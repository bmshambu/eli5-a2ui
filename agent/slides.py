"""The explanation, or something made on request, as a KPMG-branded PowerPoint deck.

Like the PDF (pdf.py), it is built on the server from the same data and sent back as a file,
because GE's frame can't download. 16:9 slides: a cover with the photo and the theme's colour,
an "at a glance" slide with the key stats, one slide per section (key points on the left; the
chart as a native, editable PowerPoint chart, the diagram as shapes, or the photo on the right;
the full text in the speaker notes), caveats and glossary, sources, and a closing slide.
A poster is one portrait slide; an email, post or one-pager is a cover and a content slide.
Pure Python (python-pptx); Arial throughout, as the brand's fallback font.
"""
import io
import re

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import brand, pdf
from .schema import KINDS, LEVELS, Chart, Creation, Diagram, Explanation

MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)
FONT = "Arial"
WHITE, DARK, MUTED = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x0C, 0x23, 0x3C), RGBColor(0x66, 0x66, 0x66)
TINT = RGBColor(0xF3, 0xF6, 0xFF)
SERIES = (brand.KPMG_BLUE, brand.PACIFIC, brand.PURPLE, brand.COBALT, brand.LIGHT_PURPLE)


def _rgb(colour) -> RGBColor:
    return RGBColor.from_string(colour.lstrip("#")) if isinstance(colour, str) else RGBColor.from_string(
        colour.hexval()[2:].upper())


def _palette(theme: str, accent: str) -> tuple[RGBColor, RGBColor, RGBColor]:
    """(main, second, highlight) for the cover: the creative accent, or the theme's cover colours."""
    if theme == "creative":
        a1, _, a3, ah = brand.ACCENTS.get(accent, brand.ACCENTS["cobalt"])
        return _rgb(a1), _rgb(a3), _rgb(ah)
    wash, _, _, hl = pdf.COVERS.get(theme, pdf.COVERS["corporate"])
    return _rgb(wash), _rgb(brand.KPMG_BLUE), _rgb(hl)


def filename(title: str) -> str:
    return pdf.filename(title)[:-4] + ".pptx"


def _box(slide, x, y, w, h, fill: RGBColor, alpha: int | None = None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if alpha is not None:  # python-pptx has no transparency setting; set it in the XML
        clr = s.fill._xPr.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha")).set("val", str(alpha * 1000))
    return s


def _text(slide, x, y, w, h, paras, size=16, colour=DARK, bold=False, bullets=False, anchor=MSO_ANCHOR.TOP,
          align=PP_ALIGN.LEFT):
    """A text box. `paras` is a string or a list; each item a paragraph (or a list of
    (text, colour) runs)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, para in enumerate([paras] if isinstance(paras, (str, tuple)) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * 0.45)
        runs = para if isinstance(para, list) else [(("■  " if bullets else "") + str(para), colour)]
        for text, c in runs:
            r = p.add_run()
            r.text = pdf._EMOJI.sub("", text)
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, c
    return tb


def _photo(slide, photo_id: str, x, y, w, h):
    """The brand photo cropped to fill the box (the photos are 5:2 banners)."""
    pic = slide.shapes.add_picture(pdf._photo_bytes(photo_id), x, y, w, h)
    want, have = w / h, 2.5
    if want < have:  # box narrower than the photo: crop the sides
        cut = (1 - want / have) / 2
        pic.crop_left = pic.crop_right = cut
    elif want > have:
        cut = (1 - have / want) / 2
        pic.crop_top = pic.crop_bottom = cut
    return pic


def _logo(slide, variant: str, x, y, h):
    slide.shapes.add_picture(str(brand.ASSETS / f"kpmg-logo-{variant}.png"), x, y, height=h)


def _title_runs(title: str, colour: RGBColor, hl: RGBColor) -> list:
    """The title with its last words in the highlight colour, as on the page and PDF."""
    words = title.split()
    k = 0 if len(words) < 3 else 2 if len(words) > 4 else 1
    runs = [(" ".join(words[:len(words) - k]), colour)]
    return runs + [(" " + " ".join(words[-k:]), hl)] if k else runs


def _title_height(title: str, size: int, width) -> Emu:
    """Roughly how tall a bold title wraps at this size, so the text under it follows closely."""
    per_line = max(int(width / Pt(size * 0.5)), 8)
    lines = -(-len(title) // per_line)
    return Emu(int(Pt(size * 1.2) * lines + Inches(0.15)))


def _footer(slide, n: int, dark: bool = False):
    colour = RGBColor(0xAC, 0xEA, 0xFF) if dark else MUTED
    _text(slide, M, H - Inches(0.45), Inches(9), Inches(0.3), pdf.LEGAL, size=8, colour=colour)
    _text(slide, W - M - Inches(3), H - Inches(0.45), Inches(3), Inches(0.3), f"KPMG. Make the Difference.   {n}",
          size=8, colour=colour, align=PP_ALIGN.RIGHT)


class _Deck:
    def __init__(self, title: str, main: RGBColor, portrait: bool = False):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = (Inches(8.27), Inches(11.69)) if portrait else (W, H)
        self.prs.core_properties.title, self.prs.core_properties.author = title, "KPMG"
        self.main, self.n = main, 0

    def slide(self):
        self.n += 1
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])  # blank

    def content(self, heading: str, number: str = "", notes: str = ""):
        """A white content slide with the heading, an accent rule, the small logo and the footer."""
        s = self.slide()
        runs = ([(number + "  ", self.main)] if number else []) + [(heading, _rgb(brand.KPMG_BLUE))]
        _text(s, M, Inches(0.45), W - 2 * M - Inches(1.4), Inches(0.9), [runs], size=28, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _box(s, M, Inches(1.35), Inches(1.2), Inches(0.08), self.main)
        _logo(s, "blue", W - M - Inches(1.1), Inches(0.55), Inches(0.43))
        _footer(s, self.n)
        if notes:
            s.notes_slide.notes_text_frame.text = pdf._EMOJI.sub("", notes)
        return s

    def save(self) -> bytes:
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def _cover(deck: _Deck, title: str, lede: str, eyebrow: str, photo: str, theme: str, accent: str, tags=()):
    main, second, hl = _palette(theme, accent)
    s = deck.slide()
    _photo(s, photo, 0, 0, W, H)
    _box(s, 0, 0, W, H, main, alpha=45)
    _box(s, 0, 0, Inches(7.6), H, main, alpha=70)
    _box(s, W - Inches(2.2), H - Inches(2.6), Inches(3.4), Inches(3.4), second, alpha=45, shape=MSO_SHAPE.OVAL)
    _logo(s, "white", M, Inches(0.5), Inches(0.55))
    _text(s, M, Inches(2.0), Inches(9), Inches(0.4), eyebrow.upper(), size=14, colour=hl, bold=True)
    th = _title_height(title, 44, Inches(10))
    _text(s, M, Inches(2.5), Inches(10), th, [_title_runs(title, WHITE, hl)], size=44, bold=True)
    _text(s, M, Inches(2.6) + th, Inches(9), Inches(1.4), lede, size=20, colour=WHITE)
    if tags:
        _text(s, M, Inches(6.1), Inches(9), Inches(0.4), "   ·   ".join(tags), size=12, colour=hl, bold=True)
    _footer(s, deck.n, dark=True)


def _stat_tiles(s, stats, x, y, w, main: RGBColor):
    stats = stats[:4]
    if not stats:
        return
    gap = Inches(0.15)
    th = Inches(1.15)
    for i, st in enumerate(stats):
        ty = y + i * (th + gap)
        _box(s, x, ty, w, th, TINT)
        _box(s, x, ty, Inches(0.08), th, _rgb(SERIES[i % 4]) if i else main)
        _text(s, x + Inches(0.25), ty + Inches(0.08), w - Inches(0.35), Inches(0.55), st.value, size=26, bold=True,
              colour=_rgb(brand.KPMG_BLUE))
        _text(s, x + Inches(0.25), ty + Inches(0.62), w - Inches(0.35), Inches(0.5), st.label, size=12)


def _chart(s, c: Chart, x, y, w, h) -> bool:
    n = len(c.labels)
    series = [se for se in c.series if len(se.values) >= n][:5]
    if not n or not series:
        return False
    data = CategoryChartData()
    data.categories = list(c.labels)
    for se in series[:1] if c.type == "donut" else series:
        data.add_series(se.name, se.values[:n])
    kind = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "donut": XL_CHART_TYPE.DOUGHNUT}[c.type]
    _text(s, x, y, w, Inches(0.4), c.title, size=14, bold=True, colour=_rgb(brand.KPMG_BLUE))
    chart = s.shapes.add_chart(kind, x, y + Inches(0.45), w, h - Inches(1.1), data).chart
    chart.font.name, chart.font.size = FONT, Pt(11)
    chart.has_title = False  # the title is above the chart already
    chart.has_legend = c.type == "donut" or len(series) > 1
    if chart.has_legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    plot = chart.plots[0]
    if c.type == "donut":
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(SERIES[i % 5])
    else:
        if c.type == "bar":
            chart.category_axis.reverse_order = True
        for i, se in enumerate(plot.series):
            fmt = se.format.line if c.type == "line" else se.format.fill
            if c.type == "line":
                fmt.color.rgb, fmt.width = _rgb(SERIES[i % 5]), Pt(2.5)
            else:
                fmt.solid()
                fmt.fore_color.rgb = _rgb(SERIES[i % 5])
        plot.has_data_labels = len(series) == 1
        if plot.has_data_labels:
            plot.data_labels.number_format = '#,##0.##' + (f'"{c.unit}"' if c.unit == "%" else "")
            plot.data_labels.number_format_is_linked = False
            plot.data_labels.font.size = Pt(10)
    foot = c.caption + (f"  Source: {c.source}" if c.source else "")
    _text(s, x, y + h - Inches(0.6), w, Inches(0.6), foot, size=11, colour=MUTED)
    return True


def _node_runs(n, main: RGBColor) -> list:
    runs = [(n.label, _rgb(brand.KPMG_BLUE))]
    if n.value:
        runs.append(("  " + n.value, main))
    return runs


def _diagram(s, d: Diagram, x, y, w, h, main: RGBColor) -> bool:
    if not d.nodes:
        return False
    if d.type == "compare":
        sides = ([n for n in d.nodes if n.side != "right"], [n for n in d.nodes if n.side == "right"])
        cw = (w - Inches(0.2)) / 2
        for i, (head, nodes) in enumerate(zip((d.left_title or "", d.right_title or ""), sides)):
            cx = x + i * (cw + Inches(0.2))
            _box(s, cx, y, cw, Inches(0.5), main if i == 0 else _rgb(brand.KPMG_BLUE))
            _text(s, cx + Inches(0.1), y, cw, Inches(0.5), head, size=14, bold=True, colour=WHITE, anchor=MSO_ANCHOR.MIDDLE)
            _box(s, cx, y + Inches(0.5), cw, h - Inches(1.1), TINT)
            paras = []
            for n in nodes:
                paras += [[(r[0], r[1]) for r in _node_runs(n, main)], n.note]
            _text(s, cx + Inches(0.15), y + Inches(0.6), cw - Inches(0.3), h - Inches(1.3), paras, size=12)
    elif d.type == "matrix" and len(d.nodes) >= 4:
        cw, ch = (w - Inches(0.1)) / 2, (h - Inches(0.8)) / 2
        for i, n in enumerate(d.nodes[:4]):
            cx, cy = x + (i % 2) * (cw + Inches(0.1)), y + (i // 2) * (ch + Inches(0.1))
            _box(s, cx, cy, cw, ch, TINT if i in (0, 3) else RGBColor(0xE9, 0xED, 0xFC))
            _text(s, cx + Inches(0.15), cy + Inches(0.1), cw - Inches(0.3), ch - Inches(0.2),
                  [_node_runs(n, main), n.note], size=12)
    else:
        nodes = d.nodes[:6]
        rh = (h - Inches(0.7)) / len(nodes)
        for i, n in enumerate(nodes):
            ry = y + i * rh
            _box(s, x, ry + Inches(0.05), Inches(0.5), Inches(0.42), _rgb(SERIES[i % 4]) if i else main)
            _text(s, x, ry + Inches(0.05), Inches(0.5), Inches(0.42), f"{i + 1:02d}", size=12, bold=True, colour=WHITE,
                  anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
            _text(s, x + Inches(0.65), ry, w - Inches(0.65), rh, [_node_runs(n, main), n.note], size=12)
    _text(s, x, y + h - Inches(0.6), w, Inches(0.6), d.caption, size=11, colour=MUTED)
    return True


def _short(paragraphs: list[str], limit: int = 3) -> list[str]:
    """Slide bullets from paragraphs: the first sentence of each, for when a section has no points."""
    return [re.split(r"(?<=[.!?])\s", p.strip(), maxsplit=1)[0] for p in paragraphs[:limit] if p.strip()]


def build(e: Explanation, sources: list[dict], level: str = "expert", documents=()) -> bytes:
    """The deck for one explanation."""
    main, _, _ = _palette(e.theme, e.accent)
    deck = _Deck(e.title, main)
    eyebrow = f'{"Document explained" if documents else "Explained"} · {LEVELS[level][0]}'
    _cover(deck, e.title, e.one_liner, eyebrow, e.hero_photo, e.theme, e.accent, e.tags[:3])

    s = deck.content("At a glance", notes=e.summary + ("\n\nThink of it like: " + e.analogy if e.analogy else ""))
    left = Inches(7.4) if e.key_stats else W - 2 * M
    _text(s, M, Inches(1.7), left, Inches(1.9), e.summary, size=15)
    _text(s, M, Inches(3.75), left, Inches(3.0), e.takeaways[:4], size=15, bullets=True, colour=_rgb(brand.KPMG_BLUE))
    _stat_tiles(s, e.key_stats, M + Inches(7.8), Inches(1.7), W - 2 * M - Inches(7.8), main)

    used = {e.hero_photo}
    for i, sec in enumerate(e.sections, 1):
        notes = "\n\n".join(sec.paragraphs) + (f"\n\nFrom the document: {sec.doc_ref}" if sec.doc_ref else "")
        s = deck.content(sec.heading, f"{i:02d}", notes)
        bullets = sec.points[:4] or _short(sec.paragraphs)
        visual_x, visual_w, top, vh = M + Inches(6.4), W - 2 * M - Inches(6.4), Inches(1.7), Inches(5.1)
        has_visual = ((sec.chart and _chart(s, sec.chart, visual_x, top, visual_w, vh))
                      or (sec.diagram and _diagram(s, sec.diagram, visual_x, top, visual_w, vh, main)))
        if not has_visual and sec.photo and sec.photo not in used:
            used.add(sec.photo)
            _photo(s, sec.photo, visual_x, top, visual_w, visual_w / 2.5)
            has_visual = True
        width = Inches(6.0) if has_visual else W - 2 * M
        if sec.doc_ref:
            _text(s, M, top, width, Inches(0.4), f"From the document: {sec.doc_ref}", size=12, colour=MUTED)
            top += Inches(0.45)
        _text(s, M, top, width, Inches(5.0), bullets, size=17, bullets=True)

    if e.caveats or e.glossary:
        s = deck.content("Caveats and glossary" if e.caveats and e.glossary else "Caveats" if e.caveats else "Glossary")
        half = (W - 2 * M - Inches(0.4)) / 2
        x = M
        if e.caveats:
            _text(s, x, Inches(1.7), half if e.glossary else W - 2 * M, Inches(5), e.caveats[:5], size=14, bullets=True)
            x += half + Inches(0.4)
        if e.glossary:
            rows = e.glossary[:8]
            t = s.shapes.add_table(len(rows) + 1, 2, x, Inches(1.7), half if e.caveats else W - 2 * M, Inches(0.4) * (len(rows) + 1)).table
            for r, (a, b) in enumerate([("Term", "Meaning")] + [(g.term, g.definition) for g in rows]):
                for col, val in enumerate((a, b)):
                    cell = t.cell(r, col)
                    cell.text = val
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.name, run.font.size = FONT, Pt(11)
                            run.font.bold = r == 0 or col == 0
                            run.font.color.rgb = WHITE if r == 0 else DARK
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(brand.KPMG_BLUE) if r == 0 else (TINT if r % 2 else WHITE)

    if documents or sources:
        s = deck.content("Based on your document" if documents else "Sources")
        lines = list(documents) + ([] if not documents else ["Explained from the document only, without web search."])
        tb = _text(s, M, Inches(1.7), W - 2 * M, Inches(5), lines or [""], size=14)
        for i, src in enumerate(sources, 1):
            p = tb.text_frame.add_paragraph()
            r = p.add_run()
            r.text = f"{i}. {src['title']}"
            r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(14), _rgb(brand.COBALT)
            r.hyperlink.address = src["uri"]

    s = deck.slide()
    _box(s, 0, 0, W, H, _rgb(brand.KPMG_BLUE))
    _logo(s, "white", M, Inches(2.6), Inches(1.0))
    _text(s, M, Inches(4.0), Inches(10), Inches(0.8), "KPMG. Make the Difference.", size=30, bold=True, colour=WHITE)
    _footer(s, deck.n, dark=True)
    return deck.save()


def _poster(c: Creation) -> bytes:
    """One portrait slide, laid out like the poster PDF."""
    main, second, hl = _palette(c.theme, c.accent)
    deck = _Deck(c.title, main, portrait=True)
    pw, ph = deck.prs.slide_width, deck.prs.slide_height
    m = Inches(0.6)
    s = deck.slide()
    _box(s, 0, 0, pw, ph, main)
    _box(s, pw - Inches(2.2), ph - Inches(4.6), Inches(4), Inches(4), second, alpha=45, shape=MSO_SHAPE.OVAL)
    photo_h = Emu(int(pw / 2.5))
    _photo(s, c.hero_photo, 0, 0, pw, photo_h)
    _logo(s, "white", m, photo_h + Inches(0.35), Inches(0.45))
    y = photo_h + Inches(1.05)
    th = _title_height(c.title, 40, pw - 2 * m)
    _text(s, m, y, pw - 2 * m, th, [_title_runs(c.title, WHITE, hl)], size=40, bold=True)
    _text(s, m, y + th, pw - 2 * m, Inches(0.9), c.tagline, size=18, colour=WHITE)
    if c.points:
        _text(s, m, y + th + Inches(0.9), pw - 2 * m, Inches(2.2), c.points[:4], size=15, colour=WHITE, bullets=True)
    if c.details:
        cw = (pw - 2 * m) / min(len(c.details), 4)
        for i, d in enumerate(c.details[:4]):
            _text(s, m + i * cw, ph - Inches(2.35), cw, Inches(0.3), d.label.upper(), size=10, bold=True, colour=hl)
            _text(s, m + i * cw, ph - Inches(2.05), cw, Inches(0.5), d.value, size=14, bold=True, colour=WHITE)
    if c.cta:
        _box(s, 0, ph - Inches(1.4), pw, Inches(0.9), WHITE)
        _text(s, m, ph - Inches(1.4), pw - 2 * m, Inches(0.9), c.cta, size=20, bold=True, colour=_rgb(brand.KPMG_BLUE),
              anchor=MSO_ANCHOR.MIDDLE)
    _text(s, m, ph - Inches(0.42), pw - 2 * m, Inches(0.3), pdf.LEGAL, size=7, colour=WHITE)
    return deck.save()


def build_creation(c: Creation) -> bytes:
    """The deck for a poster (one portrait slide), or an email, post or one-pager (cover + content)."""
    if c.kind == "poster":
        return _poster(c)
    main, _, _ = _palette(c.theme, c.accent)
    deck = _Deck(c.title, main)
    _cover(deck, c.title, c.tagline, KINDS[c.kind][0], c.hero_photo, c.theme, c.accent)
    s = deck.content(c.subject if c.kind == "email" and c.subject else c.title, notes=_plain(c))
    paras = ([c.greeting] if c.kind == "email" and c.greeting else []) + c.body[:4]
    for b in c.blocks[:3]:
        paras.append([(b.heading + ": ", _rgb(brand.KPMG_BLUE)), (b.text, DARK)])
        paras += [[("■  " + pt, DARK)] for pt in b.points[:3]]
    if c.kind == "email" and c.sign_off:
        paras += c.sign_off.splitlines()
    right = bool(c.points or c.details or c.cta)
    _text(s, M, Inches(1.7), Inches(7.2) if right else W - 2 * M, Inches(5), paras or [c.tagline], size=15)
    x, w, y = M + Inches(7.6), W - 2 * M - Inches(7.6), Inches(1.7)
    if c.points:
        _text(s, x, y, w, Inches(2.2), c.points[:5], size=14, bullets=True, colour=_rgb(brand.KPMG_BLUE))
        y += Inches(2.3)
    for d in c.details[:4]:
        _text(s, x, y, w, Inches(0.4), [[(d.label + ":  ", main), (d.value, DARK)]], size=14, bold=True)
        y += Inches(0.42)
    if c.cta:
        _box(s, x, max(y + Inches(0.2), Inches(5.6)), w, Inches(0.7), main)
        _text(s, x + Inches(0.15), max(y + Inches(0.2), Inches(5.6)), w - Inches(0.3), Inches(0.7), c.cta, size=16, bold=True,
              colour=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    if c.kind == "social" and c.hashtags:
        _text(s, M, Inches(6.5), Inches(7), Inches(0.4), " ".join("#" + h.lstrip("#") for h in c.hashtags), size=13,
              colour=_rgb(brand.COBALT), bold=True)
    return deck.save()


def _plain(c: Creation) -> str:
    from .maker import plain_text  # maker imports explainer; keep this module light to import
    return plain_text(c)
